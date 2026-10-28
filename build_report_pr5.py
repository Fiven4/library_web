import os
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_report():
    doc = docx.Document('docs/report-practice-4.docx')
    body = doc._body._element

    # 1. Update SDT Table of Contents
    sdt_updated = False
    for child in body:
        if child.tag.endswith('}sdt'):
            sdt_updated = True
            for t in child.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
                if 'ПРАКТИЧЕСКАЯ РАБОТА' in t.text:
                    t.text = 'ПРАКТИЧЕСКАЯ РАБОТА №5'
            break
    print(f"SDT TOC updated: {sdt_updated}")

    # 2. Locate Heading 1
    h1_idx = None
    for i, child in enumerate(body):
        if child.tag.endswith('}p'):
            p = docx.text.paragraph.Paragraph(child, doc)
            if 'ПРАКТИЧЕСКАЯ РАБОТА' in p.text:
                h1_idx = i
                break
    print(f"Heading 1 child index: {h1_idx}")

    # Update Heading 1
    p_h1 = docx.text.paragraph.Paragraph(body[h1_idx], doc)
    p_h1.text = "ПРАКТИЧЕСКАЯ РАБОТА №5"
    p_h1.style = doc.styles['Heading 1']
    pPr = p_h1._p.get_or_add_pPr()
    if pPr.find(qn('w:pageBreakBefore')) is None:
        pPr.append(parse_xml(f'<w:pageBreakBefore {nsdecls("w")}/>'))

    # Update Subtitle (h1_idx + 1)
    p_sub = docx.text.paragraph.Paragraph(body[h1_idx + 1], doc)
    p_sub.text = ""
    run_sub = p_sub.add_run("Аутентификация, роли и защита маршрутов")
    run_sub.bold = True
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(14)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub_pr = p_sub._p.get_or_add_pPr()
    p_sub_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:after="120" w:line="360" w:lineRule="auto"/>'))

    # Update Purpose (h1_idx + 2)
    p_purp = docx.text.paragraph.Paragraph(body[h1_idx + 2], doc)
    p_purp.text = ""
    run_purp = p_purp.add_run(
        "Цель работы: освоить проектирование и практическую реализацию подсистемы аутентификации и авторизации "
        "веб-приложения на основе протокола JWT (Access и Refresh токены), разделение прав доступа по ролевой модели "
        "(читатель, библиотекарь, администратор) и защиту маршрутов с помощью декларативной маршрутизации GoRouter. "
        "Реализовать формы входа и регистрации с динамической валидацией сложности пароля в реальном времени, "
        "автоматическое фоновое обновление access-токена при 401 Unauthorized, глобальный таймер неактивности "
        "с предупреждением и выходом из системы, сокрытие недоступных элементов интерфейса и обработку 403 Forbidden. "
        "Провести эксперимент по попытке клиентской подмены роли и доказать, что безопасность веб-приложения "
        "гарантируется исключительно серверными проверками прав доступа."
    )
    run_purp.font.name = "Times New Roman"
    run_purp.font.size = Pt(14)
    p_purp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_purp_pr = p_purp._p.get_or_add_pPr()
    p_purp_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:line="360" w:lineRule="auto"/>'))
    p_purp_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="709"/>'))

    # Update "Ход работы:" (h1_idx + 3)
    p_hod = docx.text.paragraph.Paragraph(body[h1_idx + 3], doc)
    p_hod.text = ""
    run_hod = p_hod.add_run("Ход работы:")
    run_hod.bold = False
    run_hod.font.name = "Times New Roman"
    run_hod.font.size = Pt(14)
    p_hod.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_hod_pr = p_hod._p.get_or_add_pPr()
    p_hod_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:line="360" w:lineRule="auto"/>'))
    p_hod_pr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="709"/>'))
    p_hod_pr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))

    # Remove old content after h1_idx + 3 up to sectPr
    cutoff = h1_idx + 4
    while len(body) > cutoff:
        child = body[cutoff]
        if child.tag.endswith('}sectPr'):
            break
        body.remove(child)
    print(f"Body truncated to {len(body)} elements (sectPr preserved)")

    # Helper functions
    def add_step_p(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:line="360" w:lineRule="auto"/>'))
        pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="709"/>'))
        return p

    def add_table_caption(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="120" w:after="60" w:line="360" w:lineRule="auto"/>'))
        pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:firstLine="709"/>'))
        pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
        return p

    def add_figure(img_path, caption_text, width_cm=15.5):
        if not os.path.exists(img_path):
            print(f"WARNING: Image not found: {img_path}")
            return
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img_pr = p_img._p.get_or_add_pPr()
        p_img_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="120" w:after="60"/>'))
        p_img_pr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Cm(width_cm))

        p_cap = doc.add_paragraph()
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = "Times New Roman"
        run_cap.font.size = Pt(14)
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap_pr = p_cap._p.get_or_add_pPr()
        p_cap_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="60" w:after="140" w:line="360" w:lineRule="auto"/>'))

    def add_custom_table(headers, rows, col_widths=None):
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        # Headers
        hdr_cells = table.rows[0].cells
        for col_idx, text in enumerate(headers):
            cell = hdr_cells[col_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_pr = p._p.get_or_add_pPr()
            p_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="40" w:after="40" w:line="240" w:lineRule="auto"/>'))
            run = p.add_run(text)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

        # Data rows
        for row_idx, data_row in enumerate(rows):
            row_cells = table.rows[row_idx + 1].cells
            for col_idx, val in enumerate(data_row):
                cell = row_cells[col_idx]
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
                p_pr = p._p.get_or_add_pPr()
                p_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:before="30" w:after="30" w:line="240" w:lineRule="auto"/>'))
                run = p.add_run(str(val))
                run.font.name = "Times New Roman"
                run.font.size = Pt(10.5)
                set_cell_margins(cell, top=60, bottom=60, left=100, right=100)

        if col_widths:
            for row in table.rows:
                for col_idx, width in enumerate(col_widths):
                    row.cells[col_idx].width = width

        # Spacer after table
        p_sp = doc.add_paragraph()
        p_sp_pr = p_sp._p.get_or_add_pPr()
        p_sp_pr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:after="80"/>'))

    # Step 1
    add_step_p(
        "Иерархическая ролевая модель и структура профиля пользователя. Подсистема авторизации "
        "библиотечной системы library_web строится на основе строгой ролевой модели (Role-Based Access Control, RBAC). "
        "В проекте выделено три уровня доступа: читатель (reader, уровень 1), библиотекарь (librarian, уровень 2) "
        "и администратор (admin, уровень 3). Иерархия реализована через перечисление Role в файле lib/models/role.dart, "
        "где метод canAccess(Role minRole) сопоставляет числовые уровни полномочий. Модель пользователя AppUser "
        "инкапсулирует идентификатор, имя пользователя, ФИО, адрес электронной почты, назначенную роль и опциональную "
        "ссылку на читательский билет readerId. Модель является неизменяемой (immutable) и поддерживает сериализацию "
        "в формат JSON и обратно."
    )
    add_table_caption("Таблица 1 – Матрица ролей и полномочий в библиотечной системе library_web")
    add_custom_table(
        ["Роль пользователя", "Уровень", "Доступные разделы", "Права на операции с книгами и справочниками", "Управление пользователями"],
        [
            ["reader (Читатель)", "1", "Каталог книг, Мои выдачи", "Только просмотр каталога; кнопки создания, правки и удаления скрыты", "Нет доступа (403 Forbidden)"],
            ["librarian (Библиотекарь)", "2", "Книги, Авторы, Жанры, Издательства, Читатели, Выдачи", "Полный CRUD: добавление книг, оформление выдач, возврат, управление справочниками", "Нет доступа (403 Forbidden)"],
            ["admin (Администратор)", "3", "Все разделы системы + панель управления пользователями (/admin/users)", "Полный CRUD над всеми сущностями базы данных библиотеки", "Полный доступ: просмотр списка, аудит ролей, статистика"]
        ],
        [Cm(3.5), Cm(1.8), Cm(3.5), Cm(4.5), Cm(3.2)]
    )
    add_figure("docs/figures-pr5/01-auth-arch.png", "Рисунок 1 – Архитектурная схема подсистемы аутентификации и ролевого разграничения прав доступа")
    add_figure("docs/figures-pr5/05-models-code.png", "Рисунок 2 – Иерархия ролей и модель пользователя (role.dart и app_user.dart)")

    # Step 2
    add_step_p(
        "Сетевой слой аутентификации и контракт REST API. Для взаимодействия с сервером аутентификации "
        "создан сервисный класс AuthApi в файле lib/core/auth_api.dart. Сервис взаимодействует с эндпоинтами /api/auth/* "
        "учебного сервера mock-server.js. При успешном входе или регистрации сервер генерирует пару криптографических "
        "токенов: краткоживущий Access Token со сроком действия 15 минут и долгоживущий Refresh Token со сроком жизни 7 дней. "
        "Сервис AuthApi типизирует ответы сервера в структуру AuthResult и преобразует сетевые сбои в доменные исключения "
        "UnauthorizedException, ValidationException и ServerException."
    )
    add_table_caption("Таблица 2 – Спецификация эндпоинтов аутентификации учебного REST API")
    add_custom_table(
        ["HTTP-метод", "Эндпоинт API", "Тело запроса", "Успешный ответ (код 200/204)", "Ошибки валидации и доступа"],
        [
            ["POST", "/api/auth/login", "username, password", "accessToken, refreshToken, user, expiresIn", "401 Unauthorized (неверный пароль)"],
            ["POST", "/api/auth/register", "fullName, email, username, password, readerId", "accessToken, refreshToken, user, expiresIn", "422 ValidationException (слабый пароль, дубликат email)"],
            ["POST", "/api/auth/refresh", "refreshToken", "accessToken, refreshToken, user, expiresIn", "401 Unauthorized (просроченный токен refresh)"],
            ["GET", "/api/auth/me", "Заголовок Authorization: Bearer", "user (профиль текущего пользователя)", "401 Unauthorized (отсутствие или недействительность токена)"],
            ["POST", "/api/auth/logout", "refreshToken", "HTTP 204 No Content (отзыв токена)", "–"],
            ["GET", "/api/users", "Заголовок Authorization: Bearer", "items: List<User>, total, page, size", "403 Forbidden (если роль ниже admin)"]
        ],
        [Cm(2.2), Cm(3.2), Cm(3.5), Cm(4.2), Cm(3.4)]
    )
    add_figure("docs/figures-pr5/06-auth-api-code.png", "Рисунок 3 – Реализация сетевого сервиса аутентификации AuthApi (auth_api.dart)")

    # Step 3
    add_step_p(
        "Реактивное управление состоянием сессии через AuthNotifier. Центральным элементом управления "
        "сеансом пользователя выступает реактивный провайдер AuthNotifier (ChangeNotifier) в каталоге lib/state/auth_notifier.dart. "
        "Класс обеспечивает сохранение учетных данных, маркеров доступа и меток активности в локальном хранилище "
        "SharedPreferences (localStorage веб-браузера). При запуске приложения метод restore() считывает сохраненную "
        "сессию, проверяет непревышение максимального времени сеанса (1 час) и восстанавливает контекст авторизации "
        "без повторного ввода пароля. Класс реализует методы login(), register(), refreshTokens(), logout(), "
        "а также функцию simulateClientRoleBypass() для экспериментального исследования надежности клиентской защиты."
    )
    add_figure("docs/figures-pr5/07-auth-notifier-code.png", "Рисунок 4 – Реактивный провайдер AuthNotifier с управлением токенами и сессией (auth_notifier.dart)")

    # Step 4
    add_step_p(
        "Перехватчики HTTP-клиента Dio: прозрачная инжекция токенов и Silent Refresh. Для бесшовной работы "
        "с защищенным API конфигурация Dio в ApiClient дополнена интерсепторами двух типов. Первый перехватчик onRequest "
        "считывает актуальный accessToken из AuthNotifier и автоматически прикрепляет заголовок Authorization: Bearer <token> "
        "ко всем сетевым запросам. Второй перехватчик onError анализирует статус-код ответа сервера: при получении кода "
        "401 Unauthorized и наличии сохраненного refreshToken инициируется прозрачный запрос обновления пары токенов "
        "POST /api/auth/refresh. В случае успеха оригинальный запрос повторяется с новым токеном. Чтобы исключить "
        "зацикливание при неверных реквизитах входа, перехватчик не обрабатывает запросы, содержащие путь /auth/."
    )
    add_figure("docs/figures-pr5/02-jwt-lifecycle.png", "Рисунок 5 – Жизненный цикл JWT-токенов: авторизация, прозрачное обновление при 401 и защита от зацикливания")
    add_figure("docs/figures-pr5/08-api-client-refresh-code.png", "Рисунок 6 – Перехватчик ApiClient с прозрачным обновлением токена и защитой от бесконечного цикла (api_client.dart)")

    # Step 5
    add_step_p(
        "Экран аутентификации LoginScreen и быстрое переключение учетных записей. Вход в систему "
        "выполняется через экран LoginScreen (lib/screens/login_screen.dart). Форма включает поля ввода логина и пароля "
        "с клиентской валидацией обязательности заполнения, переключатель видимости пароля и индикатор загрузки. "
        "При отказе сервера в авторизации над формой отображается акцентный баннер с текстом ошибки 401 Unauthorized. "
        "Для удобства тестирования и демонстрации преподавателю под формой размещены чипы быстрого заполнения "
        "учетных данных для трех ролей: «Читатель» (reader), «Библиотекарь» (librarian) и «Администратор» (admin). "
        "После успешной авторизации пользователь перенаправляется на исходный запрошенный маршрут, сохраненный "
        "в параметре from, либо в каталог книг по умолчанию."
    )
    add_figure("docs/figures-pr5/13-login-screen.png", "Рисунок 7 – Экран входа в систему (/login) с быстрыми демо-чипами ролей")
    add_figure("docs/figures-pr5/15-login-error-401.png", "Рисунок 8 – Обработка ошибки 401 Unauthorized при неверном логине или пароле")

    # Step 6
    add_step_p(
        "Экран регистрации RegisterScreen и валидация сложности пароля в реальном времени. Экран "
        "регистрации RegisterScreen (lib/screens/register_screen.dart) реализует создание новой учетной записи читателя. "
        "В соответствии с требованиями на повышенную оценку внедрен динамический валидатор сложности пароля, "
        "информирующий пользователя о выполнении правил непосредственно в процессе набора текста. Чек-лист правил "
        "включает три обязательных критерия: минимальную длину не менее 8 символов, наличие хотя бы одной цифры "
        "и наличие хотя бы одного специального знака. Каждый пункт сопровождается анимированной иконкой статуса, "
        "меняющей цвет с нейтрального серого на зеленый при удовлетворении условия. Кнопка отправки формы блокируется "
        "до тех пор, пока все критерии надежности пароля не будут полностью соблюдены."
    )
    add_table_caption("Таблица 3 – Требования к сложности пароля и правила валидации")
    add_custom_table(
        ["Критерий сложности", "Регулярное выражение / условие", "Индикатор интерфейса", "Статус до ввода", "Статус при выполнении"],
        [
            ["Минимальная длина (8+ символов)", "password.length >= 8", "Иконка чекбокса + текст «Не менее 8 символов»", "Серый круг (не выполнено)", "Зеленая галочка (выполнено)"],
            ["Наличие хотя бы одной цифры", "RegExp(r'[0-9]').hasMatch(p)", "Иконка чекбокса + текст «Хотя бы одна цифра»", "Серый круг (не выполнено)", "Зеленая галочка (выполнено)"],
            ["Наличие спецсимвола", "RegExp(r'[!@#$%^&*(),.?\":{}|<>]').hasMatch(p)", "Иконка чекбокса + текст «Хотя бы один спецсимвол»", "Серый круг (не выполнено)", "Зеленая галочка (выполнено)"],
            ["Подтверждение пароля", "password == confirmPassword", "Сообщение валидатора под полем повтора", "Пустое поле", "Пароли совпадают"]
        ],
        [Cm(4.0), Cm(3.8), Cm(3.8), Cm(2.4), Cm(2.5)]
    )
    add_figure("docs/figures-pr5/09-register-validator-code.png", "Рисунок 9 – Код динамического валидатора сложности пароля (register_screen.dart)")
    add_figure("docs/figures-pr5/14-register-screen.png", "Рисунок 10 – Экран регистрации (/register) с интерактивным чек-листом сложности пароля")

    # Step 7
    add_step_p(
        "Декларативная маршрутизация и защита маршрутов через GoRouter. Безопасность навигации "
        "организована на уровне маршрутизатора GoRouter (lib/core/router.dart). Провайдер AuthNotifier передается "
        "в параметр refreshListenable, благодаря чему любое изменение состояния авторизации мгновенно инициирует "
        "перерасчет глобального guard-правила redirect. При попытке неавторизованного обращения к защищенным каталогам "
        "происходит автоматический редирект на /login с сохранением целевого адреса в query-параметре from. "
        "Если аутентифицированный пользователь пытается открыть служебный раздел /admin/*, не обладая правами администратора, "
        "маршрутизатор принудительно перенаправляет запрос на экран ошибки доступа /forbidden."
    )
    add_figure("docs/figures-pr5/10-router-guard-code.png", "Рисунок 11 – Конфигурация маршрутов и ролевых guards в GoRouter (router.dart)")

    # Step 8
    add_step_p(
        "Обработка несанкционированного доступа: экран 403 Forbidden. Для корректной визуализации "
        "отказа в правах доступа разработан специализированный экран ForbiddenScreen (lib/screens/forbidden_screen.dart). "
        "Экран выводит предупреждающую пиктограмму блокировки, заголовок «Доступ запрещен (403 Forbidden)», подробное "
        "пояснение о нехватке полномочий для запрошенного действия и навигационную кнопку «Вернуться на главную», "
        "переводящую пользователя в общедоступный каталог литературы."
    )
    add_figure("docs/figures-pr5/19-forbidden-screen.png", "Рисунок 12 – Экран 403 Forbidden (/forbidden) при отказе в доступе")

    # Step 9
    add_step_p(
        "Ролевое разграничение интерфейса каталогов и действий. В интерфейсе библиотечной системы "
        "применен принцип минимизации шума: элементы управления, недоступные текущей роли, скрываются из разметки. "
        "В виджете EntityListView плавающая кнопка «Добавить» и всплывающее меню действий над строками таблицы (редактирование, "
        "удаление) отображаются исключительно для пользователей с ролью библиотекаря или администратора. Пользователь "
        "с ролью читателя видит каталог в режиме ознакомления без возможности модификации записей."
    )
    add_figure("docs/figures-pr5/16-reader-catalog.png", "Рисунок 13 – Каталог книг под ролью «Читатель»: сокрытие кнопки создания и действий")
    add_figure("docs/figures-pr5/17-librarian-catalog.png", "Рисунок 14 – Каталог книг под ролью «Библиотекарь»: полный доступ к добавлению и редактированию")

    # Step 10
    add_step_p(
        "Панель администрирования пользователей (/admin/users). Для роли администратора разработан "
        "эксклюзивный экран UsersScreen (lib/screens/users_screen.dart), доступный по маршруту /admin/users. Экран "
        "содержит блок аналитических карточек с подсчетом общего числа учетных записей и распределения по ролям, строку "
        "динамического поиска по имени и почте, а также интерактивную таблицу с бейджами ролей и датами регистрации."
    )
    add_figure("docs/figures-pr5/18-admin-users.png", "Рисунок 15 – Экран администрирования (/admin/users) со сводной статистикой и списком пользователей")

    # Step 11
    add_step_p(
        "Мониторинг неактивности пользователя и тайм-аут сеанса (InactivityWatcher). Для обеспечения "
        "безопасности публичных терминалов внедрен виджет InactivityWatcher (lib/widgets/inactivity_watcher.dart), оборачивающий "
        "все дерево приложения. Виджет отслеживает пользовательский ввод: перемещения курсора и клики мыши через "
        "HitTestBehavior.translucent, а также нажатия клавиш клавиатуры через HardwareKeyboard.instance.addHandler. "
        "При отсутствии действий в течение 2 минут 30 секунд поверх экрана открывается модальное диалоговое окно "
        "с обратным отсчетом 30 секунд. Если пользователь не подтверждает продолжение работы, по достижении 3 минут "
        "вызывается auth.logout() с отзывом сессии и принудительным выходом на экран входа."
    )
    add_figure("docs/figures-pr5/03-inactivity-flow.png", "Рисунок 16 – Схема алгоритма отслеживания неактивности пользователя и тайм-аута сеанса")
    add_figure("docs/figures-pr5/11-inactivity-watcher-code.png", "Рисунок 17 – Реализация виджета InactivityWatcher с глобальным перехватом ввода (inactivity_watcher.dart)")
    add_figure("docs/figures-pr5/20-inactivity-dialog.png", "Рисунок 18 – Предупреждающее диалоговое окно за 30 секунд до завершения сеанса по неактивности")
    add_figure("docs/figures-pr5/21-inactivity-logout.png", "Рисунок 19 – Экран входа после автоматического завершения сессии по таймауту")

    # Step 12
    add_step_p(
        "Эксперимент 17: Имитация подмены роли на клиенте и доказательство серверного рубежа безопасности. "
        "В рамках обязательного исследовательского задания проведен эксперимент по компрометации клиентского состояния. "
        "Пользователь с учетной записью читателя (reader) производит модификацию поля роли в памяти клиента "
        "(через консоль разработчика DevTools или локальное хранилище localStorage) на значение «admin». В результате "
        "реактивного уведомления виджетов в интерфейсе Flutter Web появляются элементы администратора: кнопка создания книги "
        "и вкладка «Пользователи». Однако при первой же попытке отправить сетевой запрос POST /api/books или GET /api/users "
        "сервер REST API производит валидацию криптографической подписи переданного JWT-токена, извлекает исходную роль "
        "читателя и отклоняет запрос с ошибкой 403 Forbidden («Операция доступна начиная с роли «librarian»/«admin»»). "
        "Эксперимент наглядно доказал, что любые клиентские ограничения носят исключительно характер оптимизации UX, "
        "а реальная информационная безопасность обеспечивается серверным контролем доступа."
    )
    add_figure("docs/figures-pr5/04-experiment-bypass-vs-server.png", "Рисунок 20 – Концептуальная схема Эксперимента 17: иллюзия клиентской безопасности против сервера")
    add_figure("docs/figures-pr5/22-experiment-devtools-bypass.png", "Рисунок 21 – Эксперимент 17 (Шаг 1): клиентская подмена роли на admin, появление кнопок в UI")
    add_figure("docs/figures-pr5/23-experiment-server-reject-403.png", "Рисунок 22 – Эксперимент 17 (Шаг 2): отклонение запроса сервером REST API с ошибкой 403 Forbidden")

    # Step 13
    add_step_p(
        "Автоматизированное модульное и интеграционное тестирование подсистемы авторизации. В файле "
        "test/auth_test.dart разработан комплекс из 14 специализированных автоматизированных тестов, охватывающих "
        "все уровни разработанного функционала: от иерархии прав ролей до интерсепторов сетевого клиента Dio "
        "и сквозной проверки защитных редиректов маршрутизатора GoRouter."
    )
    add_table_caption("Таблица 4 – Результаты автоматизированного тестирования подсистемы аутентификации и авторизации")
    add_custom_table(
        ["№", "Группа тестов", "Проверяемый сценарий", "Ожидаемый результат", "Статус"],
        [
            ["1", "Иерархия ролей", "Сравнение прав: reader < librarian < admin", "canAccess возвращает true только для допустимых уровней", "Успешно (Passed)"],
            ["2", "Иерархия ролей", "Сериализация и десериализация Role в строку", "Совпадение строковых имен при fromJson и toJson", "Успешно (Passed)"],
            ["3", "Валидация пароля", "Проверка минимальной длины (8 символов)", "Пароли короче 8 символов отклоняются валидатором", "Успешно (Passed)"],
            ["4", "Валидация пароля", "Проверка наличия хотя бы одной цифры", "Отклонение паролей без цифр", "Успешно (Passed)"],
            ["5", "Валидация пароля", "Проверка наличия специальных символов", "Отклонение паролей без знаков пунктуации/спецсимволов", "Успешно (Passed)"],
            ["6", "Маршрутизация", "Редирект неавторизованного с /books на /login", "GoRouter формирует адрес /login?from=%2Fbooks", "Успешно (Passed)"],
            ["7", "Маршрутизация", "Редирект читателя с /admin/users на /forbidden", "GoRouter блокирует переход и открывает 403 Forbidden", "Успешно (Passed)"],
            ["8", "Маршрутизация", "Успешный пропуск администратора на /admin/users", "Маршрут открывается без редиректа", "Успешно (Passed)"],
            ["9", "Интерсепторы Dio", "Инжекция заголовка Authorization: Bearer", "Токен корректно добавляется в headers запроса", "Успешно (Passed)"],
            ["10", "Интерсепторы Dio", "Автоматический Silent Refresh при 401", "Перехват 401, вызов /auth/refresh и повтор запроса", "Успешно (Passed)"],
            ["11", "Интерсепторы Dio", "Предотвращение бесконечного цикла на /auth/*", "Ошибка 401 на /auth/login пробрасывается дальше", "Успешно (Passed)"],
            ["12", "Сессия пользователя", "Таймер неактивности и сброс по активности", "Вызов recordActivity() обновляет метку времени", "Успешно (Passed)"],
            ["13", "Сессия пользователя", "Очистка токенов при принудительном logout", "SharedPreferences очищается, статус isAuthenticated = false", "Успешно (Passed)"],
            ["14", "Эксперимент 17", "Подмена роли на клиенте и отказ сервера 403", "Сервер возвращает 403 Forbidden, доказывая надежность", "Успешно (Passed)"]
        ],
        [Cm(0.8), Cm(2.8), Cm(4.2), Cm(5.5), Cm(2.2)]
    )
    add_figure("docs/figures-pr5/12-unit-tests-code.png", "Рисунок 23 – Код модульных тестов ролевой защиты и эксперимента 17 (auth_test.dart)")

    # Step 14
    add_step_p(
        "Итоговая верификация: статический анализ и выполнение полного пакета тестов. Перед подготовкой "
        "отчета проведена полная верификация исходного кода. Команда flutter analyze подтвердила отсутствие предупреждений "
        "и ошибок статической типизации (0 issues found). Запуск полного набора автоматизированных тестов проекта командой "
        "flutter test завершился успешным прохождением всех 152 тестов (152/152 passed) без единого сбоя."
    )
    add_figure("docs/figures-pr5/24-test-run-console.png", "Рисунок 24 – Вывод консоли терминала: flutter analyze без ошибок и 152 пройденных теста")

    # Результат работы
    add_step_p("Результат работы:")
    add_figure("docs/figures-pr5/25-result-dashboard.png", "Рисунок 25 – Результат работы: общий вид системы под учётной записью администратора с активным дашбордом")
    add_figure("docs/figures-pr5/26-result-mobile-drawer.png", "Рисунок 26 – Результат работы: адаптивный интерфейс с мобильной навигацией и профилем пользователя")

    # Вывод
    add_step_p(
        "Вывод: в ходе выполнения практической работы №5 клиент библиотечной системы library_web успешно оснащен "
        "полнофункциональной подсистемой аутентификации, авторизации и защиты маршрутов, удовлетворяющей всем требованиям "
        "на максимальную оценку «5». Реализована иерархическая трехзвенная ролевая модель (читатель, библиотекарь, администратор) "
        "с динамической адаптацией интерфейса и сокрытием недоступных операций. Внедрен полноценный жизненный цикл JWT-токенов "
        "с прозрачным обновлением access-токена через refresh-токен при коде 401 Unauthorized и гарантированной защитой от бесконечных циклов. "
        "Разработана декларативная защита маршрутов GoRouter с перенаправлением неавторизованных пользователей на форму входа "
        "с сохранением целевого адреса в параметре from, а также обработка несанкционированного доступа с экраном 403 Forbidden. "
        "Обеспечена безопасность пользовательских сеансов через таймер неактивности с предупреждением за 30 секунд и автовыходом через 3 минуты. "
        "Проведен эксперимент 17, наглядно доказавший приоритет серверного рубежа безопасности над клиентскими проверками разметки. "
        "Качество и надежность разработанных решений подтверждены успешным выполнением статического анализатора flutter analyze "
        "и 100% прохождением всех 152 автоматизированных тестов проекта."
    )

    output_path = 'docs/report-practice-5.docx'
    doc.save(output_path)
    print(f"Report successfully built and saved to {output_path}!")

if __name__ == '__main__':
    build_report()